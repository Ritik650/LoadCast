"""
PLAYHACK / ML  -  Render the HTML deck to PDF and PPTX.

    python src/export_deck.py

PDF  : Chromium print-to-PDF at exactly 13.333 x 7.5 in (16:9), one slide/page.
PPTX : one full-bleed 2x-DPI image per slide, so the layout is pixel-identical
       to the deck. Requires playwright (with Chrome) and python-pptx.
"""
import os
import asyncio
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK = os.path.join(ROOT, "deck")
HTML = os.path.join(DECK, "loadcast.html")
SHOTS = os.path.join(DECK, "_slides")
W_IN, H_IN = 13.3333, 7.5


async def render():
    os.makedirs(SHOTS, exist_ok=True)
    async with async_playwright() as pw:
        b = await pw.chromium.launch(channel="chrome")

        # ---- PDF: print CSS, exact 16:9 page, backgrounds on
        pg = await b.new_page(viewport={"width": 1400, "height": 900})
        await pg.goto("file:///" + HTML.replace("\\", "/"))
        await pg.wait_for_timeout(3000)
        pdf = os.path.join(DECK, "Loadcast_PlayHack_ML.pdf")
        await pg.pdf(path=pdf, width=f"{W_IN}in", height=f"{H_IN}in",
                     print_background=True, prefer_css_page_size=False,
                     margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
        print("PDF  -> %s (%.1f MB)" % (pdf, os.path.getsize(pdf) / 1e6))

        # ---- slide images at 2x for the PPTX
        pg2 = await b.new_page(viewport={"width": 1280, "height": 720},
                               device_scale_factor=2)
        await pg2.goto("file:///" + HTML.replace("\\", "/"))
        await pg2.wait_for_timeout(3000)
        n = await pg2.locator("section.slide").count()
        paths = []
        for i in range(n):
            f = os.path.join(SHOTS, "%02d.png" % i)
            await pg2.locator("section.slide").nth(i).screenshot(path=f)
            paths.append(f)
        print("PNG  -> %d slides at 2x" % n)
        await b.close()
        return paths


def build_pptx(paths):
    prs = Presentation()
    prs.slide_width = Inches(W_IN)
    prs.slide_height = Inches(H_IN)
    blank = prs.slide_layouts[6]
    for f in paths:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(f, 0, 0, width=prs.slide_width,
                             height=prs.slide_height)
    out = os.path.join(DECK, "Loadcast_PlayHack_ML.pptx")
    prs.save(out)
    print("PPTX -> %s (%.1f MB, %d slides)"
          % (out, os.path.getsize(out) / 1e6, len(paths)))


if __name__ == "__main__":
    build_pptx(asyncio.run(render()))
